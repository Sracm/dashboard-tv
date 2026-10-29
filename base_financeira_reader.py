# -*- coding: utf-8 -*-
import os
import openpyxl
from datetime import datetime

CAMINHO_BASE_FINANCEIRA = r"C:\Users\amello\OneDrive - MQHAIR\T.I\PROJETOS BI\DIRETORIA GERAL COMERCIAL\BASE FINANCEIRA.xlsx"

_CACHE_FINANCEIRA = None
_CACHE_MTIME = 0

def obter_dados_base_financeira():
    global _CACHE_FINANCEIRA, _CACHE_MTIME
    if not os.path.exists(CAMINHO_BASE_FINANCEIRA):
        return _CACHE_FINANCEIRA or {}

    try:
        mtime = os.path.getmtime(CAMINHO_BASE_FINANCEIRA)
        if _CACHE_FINANCEIRA is not None and mtime == _CACHE_MTIME:
            return _CACHE_FINANCEIRA

        wb = openpyxl.load_workbook(CAMINHO_BASE_FINANCEIRA, data_only=True)
        if "BASE FINANC." not in wb.sheetnames:
            return _CACHE_FINANCEIRA or {}
            
        ws = wb["BASE FINANC."]
        data = {}

        for r in range(2, ws.max_row + 1):
            ano = ws.cell(r, 1).value
            trim = ws.cell(r, 2).value
            seq = ws.cell(r, 3).value
            tipo = ws.cell(r, 4).value
            orc = ws.cell(r, 5).value
            real = ws.cell(r, 7).value
            pct = ws.cell(r, 9).value

            if ano and trim and tipo:
                ano = int(ano)
                trim_str = str(trim).strip()
                t_num = trim_str[0]
                t_key = f"{t_num}T"
                
                if ano not in data:
                    data[ano] = {}
                if t_key not in data[ano]:
                    data[ano][t_key] = {
                        "trimestre_num": int(t_num),
                        "label": f"{t_num}º TRIMESTRE",
                        "itens": []
                    }
                
                val_pct = 0.0
                is_plus_100 = False
                if isinstance(pct, (int, float)):
                    val_pct = round(pct * 100, 2)
                elif isinstance(pct, str) and "+" in pct:
                    val_pct = 100.0
                    is_plus_100 = True
                
                tipo_limpo = str(tipo).strip().upper()
                
                data[ano][t_key]["itens"].append({
                    "sequencia": seq or 99,
                    "tipo": tipo_limpo,
                    "percentual": val_pct,
                    "is_plus_100": is_plus_100,
                    "percentual_display": "+100%" if is_plus_100 else f"{val_pct:.2f}%".replace(".", ","),
                    "orcado": orc,
                    "realizado": real
                })

        for a in data:
            for t in data[a]:
                data[a][t]["itens"].sort(key=lambda x: x["sequencia"])

        _CACHE_FINANCEIRA = data
        _CACHE_MTIME = mtime
        return data

    except Exception as e:
        print(f"[BASE FINANCEIRA] Erro ao carregar planilha: {e}")
        return _CACHE_FINANCEIRA or {}